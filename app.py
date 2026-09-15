import io
import re

import pandas as pd
import streamlit as st
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

st.set_page_config(page_title="Unificar resultados · 2.º ESO", page_icon="📊", layout="wide")

st.title("Unificar resultados · Prueba inicial 2.º ESO")
st.write("Sube los Excel de toda la clase y la aplicación los reunirá en un único archivo. También puedes introducir manualmente los datos de alumnos que no tengan el Excel disponible.")

AREAS = ["Comprensión", "Morfología", "Semántica", "Textos", "Literatura", "Sintaxis"]
AREA_KEYS = {
    "Comprensión": "comprension", "Morfología": "morfologia", "Semántica": "semantica",
    "Textos": "textos", "Literatura": "literatura", "Sintaxis": "sintaxis"
}


def norm(s):
    s = str(s).strip().lower()
    s = s.replace("á", "a").replace("é", "e").replace("í", "i").replace("ó", "o").replace("ú", "u").replace("ñ", "n")
    return re.sub(r"[^a-z0-9]+", "", s)


def number(v):
    if pd.isna(v) or str(v).strip() == "":
        return None
    s = str(v).strip().replace("%", "").replace(",", ".")
    m = re.search(r"-?\d+(?:\.\d+)?", s)
    return float(m.group()) if m else None


def find_col(df, candidates):
    wanted = {norm(x) for x in candidates}
    for c in df.columns:
        if norm(c) in wanted:
            return c
    return None


def extract_from_table(df, filename):
    df = df.copy()
    df.columns = [str(c).strip() for c in df.columns]
    out = []
    name_col = find_col(df, ["name", "nombre", "alumno", "nombre y apellidos", "nombreyapellidos"])
    group_col = find_col(df, ["group", "grupo"])
    date_col = find_col(df, ["date", "fecha", "fecha y hora", "fechayhora"])
    if name_col:
        for _, r in df.iterrows():
            name = str(r.get(name_col, "")).strip()
            if not name or name.lower() in {"nan", "none"}:
                continue
            item = {
                "Alumno": name,
                "Grupo": r.get(group_col, "") if group_col else "",
                "Fecha": r.get(date_col, "") if date_col else "",
                "Fuente": filename,
            }
            for area, key in AREA_KEYS.items():
                col = find_col(df, [key, area])
                item[area] = number(r.get(col)) if col else None
            for label, candidates in {
                "Nota sobre 9": ["nota_final_sobre_9", "nota de esta parte (sobre 9)", "nota sobre 9"],
                "Ortografía": ["faltas_ortografia", "faltas de ortografía detectadas"],
                "Tildes": ["faltas_tildes", "faltas de tilde detectadas"],
            }.items():
                col = find_col(df, candidates)
                item[label] = number(r.get(col)) if col else None
            out.append(item)
    return out


def extract_vertical(df, filename):
    rows = []
    for _, r in df.iterrows():
        vals = [str(x).strip() if not pd.isna(x) else "" for x in r.tolist()]
        if vals:
            rows.append(vals)
    data = {}
    for vals in rows:
        if len(vals) >= 2:
            data[norm(vals[0])] = vals[1]

    name = data.get("alumno") or data.get("nombre") or data.get("nombreyapellidos")
    if not name:
        return []

    item = {
        "Alumno": name.strip(),
        "Grupo": data.get("grupo", ""),
        "Fecha": data.get("fechayhora", data.get("fecha", "")),
        "Fuente": filename,
    }
    for area, key in AREA_KEYS.items():
        val = data.get(key, data.get(norm(area)))
        item[area] = number(val)

    for label, keys in {
        "Nota sobre 9": ["notadeestapartesobre9", "notasobre9", "notafinalsobre9"],
        "Ortografía": ["faltasdeortografiadetectadas", "faltasortografia"],
        "Tildes": ["faltasdetildedetectadas", "faltastildes"],
    }.items():
        item[label] = number(next((data[k] for k in keys if k in data), None))
    return [item]


def read_excel(upload):
    sheets = pd.read_excel(upload, sheet_name=None, header=0)
    records = []
    for _, df in sheets.items():
        table_records = extract_from_table(df, upload.name)
        if table_records:
            records.extend(table_records)
        else:
            records.extend(extract_vertical(df, upload.name))
    return records


def sort_by_first_surname(df):
    """Ordena por el primer apellido, suponiendo el formato habitual: Nombre Apellido1 Apellido2."""
    if df.empty or "Alumno" not in df.columns:
        return df.copy()

    def surname_key(name):
        parts = str(name).strip().split()
        if len(parts) >= 2:
            return parts[1].lower()
        return str(name).lower()

    return (
        df.assign(_orden=df["Alumno"].map(surname_key))
        .sort_values(["_orden", "Alumno"], kind="stable")
        .drop(columns="_orden")
        .reset_index(drop=True)
    )


def prepare_students(df):
    df = df.copy()
    for c in AREAS + ["Nota sobre 9", "Ortografía", "Tildes"]:
        if c not in df.columns:
            df[c] = None
    df = sort_by_first_surname(df)
    df["Producción escrita"] = None
    df["Descuento producción"] = None
    df["Nota producción escrita final"] = None
    df["Nota final sobre 10"] = None
    return df[["Alumno", "Grupo", "Fecha"] + AREAS + ["Nota sobre 9", "Producción escrita", "Descuento producción", "Nota final sobre 10", "Ortografía", "Tildes", "Nota producción escrita final", "Fuente"]]


def group_label(group):
    return re.sub(r"\s+", "", str(group).strip())


def get_groups(df):
    if df.empty or "Grupo" not in df.columns:
        return []
    groups = [str(x).strip() for x in df["Grupo"].dropna().unique() if str(x).strip()]
    return sorted(groups, key=lambda x: group_label(x).lower())


def classify_student(row):
    """Clasifica al alumno y genera una observación personalizada a partir de sus áreas."""
    values = pd.to_numeric(pd.Series({area: row.get(area) for area in AREAS}), errors="coerce").dropna()
    if values.empty:
        return "SIN SEGUIMIENTO", "No hay datos suficientes de las áreas evaluadas para establecer dificultades concretas."

    deficient_areas = values[values < 5].sort_values()
    n_deficient = len(deficient_areas)
    mean = values.mean()
    areas_text = ", ".join(deficient_areas.index.tolist())

    # DESDOBLE: dificultades muy extendidas y perfil global comprometido.
    if n_deficient >= 4 or (n_deficient >= 3 and mean < 5):
        if n_deficient >= 4:
            return "DESDOBLE", f"Dificultades muy graves y generalizadas en prácticamente todas las áreas; los resultados más bajos aparecen en {areas_text}."
        return "DESDOBLE", f"Dificultades muy importantes y generalizadas, especialmente en {areas_text}."

    # REFUERZO: dificultades importantes, pero más localizadas.
    if n_deficient >= 2 and mean < 6:
        return "REFUERZO", f"Dificultades claras en {areas_text}; el resto del perfil permite seguir el aula con apoyo."
    if n_deficient == 1 and float(deficient_areas.iloc[0]) < 4.5:
        return "REFUERZO", f"Dificultad destacada en {areas_text}; el resto del perfil permite seguir el aula con apoyo."

    # SIN SEGUIMIENTO: rendimiento razonablemente funcional, aunque pueda haber dificultades puntuales.
    if n_deficient == 0:
        return "SIN SEGUIMIENTO", "Buen rendimiento global; no presenta dificultades destacadas en las áreas evaluadas."
    if n_deficient == 1:
        return "SIN SEGUIMIENTO", f"Rendimiento global adecuado; la dificultad principal es {areas_text}."
    return "SIN SEGUIMIENTO", f"Perfil relativamente equilibrado; algunas dificultades en {areas_text}, pero puede seguir con apoyo ordinario."


def build_seguimiento(final_result):
    clasificaciones = final_result.apply(classify_student, axis=1, result_type="expand")
    seguimiento = final_result[["Alumno", "Nota final sobre 10"]].copy()
    seguimiento = seguimiento.rename(columns={"Nota final sobre 10": "Nota final"})
    seguimiento["Categoría"] = clasificaciones[0].values
    seguimiento["Observaciones"] = clasificaciones[1].values
    seguimiento = seguimiento[["Alumno", "Nota final", "Categoría", "Observaciones"]]
    return sort_by_first_surname(seguimiento)


def format_excel(writer):
    header_fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
    header_font = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center")
    widths = {
        "Resultados": [34.57, 15.29, 20.29, 21.29, 19.43, 18.86, 15.43, 18.29, 16.43, 20.86, 25.71, 29.14, 26.14, 15, 11.29, 34.43],
        "Resultado final": [34.57, 15.29, 26.14],
        "Para refuerzodesdoble": [34.57, 18.57, 19.29, 88.86],
        "Medias por áreas": [28, 15],
        "Áreas": [34.57, 15.29, 21.29, 19.43, 18.86, 15.43, 18.29, 16.43, 19.43, 25.71],
    }
    for ws in writer.book.worksheets:
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center
        ws.freeze_panes = "B2" if not ws.title.startswith("Medias por áreas") else None
        base_title = next((key for key in widths if ws.title.startswith(key)), None)
        for i, width in enumerate(widths.get(base_title, []), start=1):
            ws.column_dimensions[get_column_letter(i)].width = width
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "0.00"
        if ws.title.startswith("Resultados "):
            for row in ws.iter_rows(min_row=2, min_col=4, max_col=16):
                for cell in row:
                    cell.number_format = "0.00"
        elif ws.title.startswith("Resultado final "):
            for row in ws.iter_rows(min_row=2, min_col=3, max_col=3):
                row[0].number_format = "0.00"
        elif ws.title.startswith("Para refuerzodesdoble "):
            for row in ws.iter_rows(min_row=2, min_col=2, max_col=2):
                row[0].number_format = "0.00"
        elif ws.title.startswith("Medias por áreas"):
            ws.column_dimensions["A"].width = 28
            ws.column_dimensions["B"].width = 15
            for cell in ws[1]:
                cell.alignment = center
            for row in ws.iter_rows(min_row=2, min_col=2, max_col=2):
                row[0].number_format = "0.00"
        elif ws.title.startswith("Áreas"):
            for row in ws.iter_rows(min_row=2, min_col=3, max_col=10):
                for cell in row:
                    cell.number_format = "0.00"


uploads = st.file_uploader("Sube los Excel de los alumnos", type=["xlsx", "xls"], accept_multiple_files=True)

st.markdown(
    """
    <div style="border: 2px solid #1976d2; border-radius: 8px; padding: 12px 16px; margin: 18px 0 10px 0; background-color: #f4f8ff;">
        <div style="color: #1565c0; font-size: 1.15rem; font-weight: 700;">AÑADIR ALUMNOS MANUALMENTE</div>
        <div style="color: #333; margin-top: 5px;">Utiliza este apartado para alumnos que hayan realizado la prueba en papel, que solo tengan una captura de los resultados o cuyo Excel no esté disponible. Puedes introducir sus datos aquí aunque no hayas subido ningún Excel.</div>
    </div>
    """,
    unsafe_allow_html=True,
)

manual_template = pd.DataFrame(
    [{"Alumno": "", "Grupo": "", **{area: None for area in AREAS}, "Nota sobre 9": None, "Producción escrita": None, "Descuento producción": None} for _ in range(5)]
)
manual_edit = st.data_editor(
    manual_template,
    hide_index=True,
    use_container_width=True,
    key="alumnos_manuales",
    num_rows="dynamic",
    column_config={
        "Alumno": st.column_config.TextColumn("Alumno"),
        "Grupo": st.column_config.TextColumn("Grupo"),
        **{area: st.column_config.NumberColumn(area, min_value=0.0, max_value=10.0, step=0.01, format="%.2f") for area in AREAS},
        "Nota sobre 9": st.column_config.NumberColumn("Nota sobre 9", min_value=0.0, max_value=9.0, step=0.01, format="%.2f"),
        "Producción escrita": st.column_config.NumberColumn("Producción escrita (0–1)", min_value=0.0, max_value=1.0, step=0.01, format="%.2f"),
        "Descuento producción": st.column_config.NumberColumn("Descuento por faltas", min_value=-2.0, max_value=0.0, step=0.01, format="%.2f"),
    },
)

manual_edit = manual_edit[manual_edit["Alumno"].fillna("").astype(str).str.strip() != ""].copy()
manual_edit["Fecha"] = "Manual"
manual_edit["Ortografía"] = None
manual_edit["Tildes"] = None
manual_edit["Fuente"] = "Entrada manual"

all_records = []
errors = []
for upload in uploads or []:
    try:
        all_records.extend(read_excel(upload))
    except Exception as e:
        errors.append(f"{upload.name}: {e}")

if errors:
    st.warning("Algunos archivos no se han podido leer: " + " | ".join(errors))

excel_df = prepare_students(pd.DataFrame(all_records)) if all_records else pd.DataFrame()

if not excel_df.empty or not manual_edit.empty:
    if not manual_edit.empty:
        manual_df = prepare_students(manual_edit)
    else:
        manual_df = pd.DataFrame(columns=excel_df.columns if not excel_df.empty else ["Alumno", "Grupo", "Fecha"] + AREAS + ["Nota sobre 9", "Producción escrita", "Descuento producción", "Nota final sobre 10", "Ortografía", "Tildes", "Nota producción escrita final", "Fuente"])

    df = pd.concat([excel_df, manual_df], ignore_index=True)
    df = sort_by_first_surname(df)
    st.success(f"Se han encontrado {len(df)} resultados.")

    st.markdown(
        """
        <div style="border: 2px solid #d32f2f; border-radius: 8px; padding: 12px 16px; margin: 10px 0 18px 0; background-color: #fff5f5;">
            <div style="color: #c62828; font-size: 1.15rem; font-weight: 700;">IMPORTANTE: INTRODUCE AQUÍ LOS DATOS DE PRODUCCIÓN ESCRITA</div>
            <div style="color: #333; margin-top: 5px;">Introduce la nota de producción escrita (0–1) y, en la columna de descuento, escribe directamente lo que hay que restar por faltas. Se pueden introducir descuentos de hasta 2 puntos, por ejemplo, -0,25 o -2,00. La aplicación calculará automáticamente la nota final.</div>
        </div>
        """,
        unsafe_allow_html=True,
    )

    edit_cols = ["Alumno", "Grupo"] + AREAS + ["Nota sobre 9", "Producción escrita", "Descuento producción"]
    edited = st.data_editor(
        df[edit_cols],
        hide_index=True,
        use_container_width=True,
        key="datos_produccion",
        column_config={
            "Producción escrita": st.column_config.NumberColumn("Producción escrita (0–1)", min_value=0.0, max_value=1.0, step=0.05, format="%.2f"),
            "Descuento producción": st.column_config.NumberColumn("Descuento por faltas", min_value=-2.0, max_value=0.0, step=0.05, format="%.2f"),
        },
    )

    edited["Descuento producción"] = edited["Descuento producción"].fillna(0).round(2)
    edited["Nota producción escrita final"] = (edited["Producción escrita"].fillna(0) + edited["Descuento producción"]).clip(lower=0, upper=1).round(2)
    edited["Nota final sobre 10"] = (edited["Nota sobre 9"].fillna(0) + edited["Nota producción escrita final"].fillna(0)).clip(upper=10).round(2)

    st.subheader("Resultado final")
    result_cols = ["Alumno", "Grupo", "Nota final sobre 10"]
    final_result = edited[result_cols].copy()
    final_result = sort_by_first_surname(final_result)
    st.dataframe(final_result, hide_index=True, use_container_width=True)

    st.subheader("Media de la clase por apartados")
    chart_df = edited[AREAS].mean().rename("Media").to_frame()
    chart_df.loc["Producción escrita"] = edited["Nota producción escrita final"].mean() * 10
    st.bar_chart(chart_df, y="Media")
    st.caption("El gráfico incluye las áreas del examen y la producción escrita. La producción se muestra sobre 10 para hacerla comparable visualmente.")

    final_df = df.copy()
    final_df["Producción escrita"] = edited["Producción escrita"]
    final_df["Descuento producción"] = edited["Descuento producción"]
    final_df["Nota producción escrita final"] = edited["Nota producción escrita final"]
    final_df["Nota final sobre 10"] = edited["Nota final sobre 10"]
    final_df = final_df.drop(columns=["Fuente"])
    final_df = final_df[["Alumno", "Grupo", "Fecha"] + AREAS + ["Nota sobre 9", "Producción escrita", "Descuento producción", "Nota final sobre 10", "Ortografía", "Tildes", "Nota producción escrita final"]]
    final_df = sort_by_first_surname(final_df)

    final_areas = edited[["Alumno", "Grupo"] + AREAS + ["Producción escrita"]].copy()
    final_areas = sort_by_first_surname(final_areas)

    groups = get_groups(df)
    if not groups:
        groups = [""]

    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        for group in groups:
            suffix = group_label(group) if group else "2ºA"
            mask = df["Grupo"].fillna("").astype(str).str.strip() == group
            group_names = df.loc[mask, "Alumno"].tolist()
            group_mask_edited = edited["Alumno"].isin(group_names)

            group_final_df = final_df[final_df["Alumno"].isin(group_names)].copy()
            group_final_result = final_result[final_result["Alumno"].isin(group_names)].copy()
            group_final_areas = final_areas[final_areas["Alumno"].isin(group_names)].copy()

            group_seguimiento = build_seguimiento(edited[group_mask_edited][["Alumno", "Nota final sobre 10"] + AREAS].copy())

            group_chart_df = edited[group_mask_edited][AREAS].mean().rename("Media").to_frame()
            group_chart_df.loc["Producción escrita"] = edited.loc[group_mask_edited, "Nota producción escrita final"].mean() * 10

            group_final_df.to_excel(writer, sheet_name=f"Resultados {suffix}", index=False)
            group_final_result.to_excel(writer, sheet_name=f"Resultado final {suffix}", index=False)
            group_seguimiento.to_excel(writer, sheet_name=f"Para refuerzodesdoble {suffix}", index=False)
            group_chart_df.round(2).to_excel(writer, sheet_name=f"Medias por áreas {suffix}")
            group_final_areas.to_excel(writer, sheet_name=f"Áreas {suffix}", index=False)

        format_excel(writer)

    st.markdown(
        """
        <style>
        div.stDownloadButton > button {
            background-color: #d32f2f;
            color: white;
            border: 2px solid #b71c1c;
            font-weight: 700;
            font-size: 1.05rem;
            padding: 0.65rem 1rem;
        }
        div.stDownloadButton > button:hover {
            background-color: #b71c1c;
            color: white;
            border-color: #8f1515;
        }
        </style>
        """,
        unsafe_allow_html=True,
    )

    st.download_button("DESCARGAR EXCEL DE LA CLASE", buffer.getvalue(), file_name="Resultados_unificados_2ESO_Prueba_inicial.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", use_container_width=True)
